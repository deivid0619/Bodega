"""Los Excel que se descargan: el inventario (varias hojas), lo que hay que
pedir y el historial. Archivos .xlsx de verdad: encabezados con los colores
de Pigmalion, filtros, columnas fijas, totales que respetan el filtro,
colores por estado y listos para imprimir."""
from __future__ import annotations

import io
import re
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy.orm import Session

from app import models
from app.core.config import settings
from app.modules.bodega.logic import DISPATCH
from app.modules.catalogo import service as catalog
from app.modules.inventario import serializers as ser, service as inv

INK, LIME, LIME_DEEP, FOG, MUTED = "0B0B0B", "C0FF00", "6A9A00", "F3F6EF", "6B6F68"
WHITE = "FFFFFF"
INT, MONEY, DATE = "#,##0", '"$" #,##0', "dd/mm/yyyy hh:mm"
SIZE_ORDER = ["XXS", "XS", "S", "M", "L", "XL", "XXL", "2XL", "XXXL", "3XL", "4XL", "5XL"]
STATUS = {  # estado -> (fondo, letra)
    "Agotado": ("FDE2E1", "B42318"),
    "Bajo mínimo": ("FEF0C7", "93370D"),
    "OK": ("E3F7C2", "3B5B00"),
}
MONTHS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
          "noviembre", "diciembre"]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def size_rank(size: str) -> float:
    s = (size or "").upper()
    if s in SIZE_ORDER:
        return SIZE_ORDER.index(s)
    try:
        return 100 + float(s)
    except ValueError:
        return 500


def _when(now: datetime) -> str:
    hour = now.strftime("%I:%M").lstrip("0")
    return f"{now.day} de {MONTHS[now.month - 1]} de {now.year}, {hour} {'a. m.' if now.hour < 12 else 'p. m.'}"


def _store_url(handle: str) -> str | None:
    """El enlace a la prenda en la tienda (solo si el catalogo trae su handle)."""
    base = re.sub(r"/products\.json.*$", "", settings.catalog_url or "")
    return f"{base}/products/{handle}" if base and handle and re.fullmatch(r"[a-z0-9][a-z0-9-]*", handle) else None


# ------------------------------------------------------------------ estilo
def _title(ws, title: str, subtitle: str, width: int) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color=INK)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Calibri", size=10, color=MUTED)
    for c in range(1, width + 1):  # la linea verde de Pigmalion bajo el titulo
        ws.cell(row=2, column=c).border = Border(bottom=Side(style="medium", color=LIME))
    ws.row_dimensions[1].height = 28


def _table(ws, name: str, header_row: int, headers: list[str], rows: list[list], formats: dict[str, str] | None = None,
           stripes: bool = True, band_by: int | None = None, totals: list[str] | None = None) -> int:
    """Una tabla de Excel (con filtros) desde header_row. band_by: la columna
    (indice) que agrupa filas con el mismo fondo (ej. todas las tallas de una
    referencia). totals: columnas a sumar abajo (respetan el filtro).
    Devuelve la ultima fila con datos."""
    formats = formats or {}
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30
    band, last_key, shade = PatternFill("solid", fgColor=FOG), object(), False
    for r, row in enumerate(rows, header_row + 1):
        if band_by is not None and row[band_by] != last_key:
            shade, last_key = not shade, row[band_by]
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            fmt = formats.get(headers[c - 1])
            if fmt:
                cell.number_format = fmt
            if band_by is not None and shade:
                cell.fill = band
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
    last = header_row + max(len(rows), 1)
    if not rows:  # una tabla de Excel necesita al menos una fila
        ws.cell(row=last, column=1, value="(vacío)").font = Font(italic=True, color=MUTED)
    tab = Table(displayName=name, ref=f"A{header_row}:{get_column_letter(len(headers))}{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=stripes and band_by is None)
    ws.add_table(tab)
    if totals and rows:
        tr = last + 1
        ws.cell(row=tr, column=1, value="Total (lo que se ve con el filtro)").font = Font(bold=True)
        for h in totals:
            c = headers.index(h) + 1
            col = get_column_letter(c)
            cell = ws.cell(row=tr, column=c, value=f"=SUBTOTAL(109,{col}{header_row + 1}:{col}{last})")
            cell.number_format = formats.get(h, INT)
            cell.font = Font(bold=True)
        for c in range(1, len(headers) + 1):
            ws.cell(row=tr, column=c).border = Border(top=Side(style="medium", color=INK))
    _widths(ws, headers, rows)
    return last


def _widths(ws, headers: list[str], rows: list[list], cap: int = 48) -> None:
    for c, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [len(f"{r[c - 1]:,}" if isinstance(r[c - 1], (int, float)) else str(r[c - 1] or ""))
                                         for r in rows[:2000]])
        ws.column_dimensions[get_column_letter(c)].width = min(cap, max(8, longest * 1.1 + 2))


def _print(ws, header_row: int, landscape: bool = True) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.oddFooter.center.text = "Bodega Pigmalion · página &P de &N"


def _status_colors(ws, col: str, first: int, last: int) -> None:
    for text, (bg, fg) in STATUS.items():
        ws.conditional_formatting.add(f"{col}{first}:{col}{last}", CellIsRule(
            operator="equal", formula=[f'"{text}"'], fill=PatternFill("solid", fgColor=bg), font=Font(color=fg, bold=True)))


def _bars(ws, col: str, first: int, last: int, color: str = LIME_DEEP) -> None:
    ws.conditional_formatting.add(f"{col}{first}:{col}{last}", DataBarRule(
        start_type="num", start_value=0, end_type="max", color=color, showValue=True))


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ datos
def _inventory_rows(db: Session) -> dict:
    """Todo lo del inventario, una vez: por talla, con bodega, reserva, de
    paso y outlet separados (como en la app)."""
    names = ser.loc_names(db)
    stock = ser.stock_map(db)
    outs = ser.out_30d_map(db)
    outlet = inv.outlet_ids(db)
    index = inv._reserve_index(db)
    cat = catalog.items()  # los precios y enlaces de la tienda (si hace rato no se consultan, se traen)
    products = sorted(db.query(models.Product).all(), key=lambda p: (p.name, size_rank(p.size), p.sku))
    rows = []
    for p in products:
        places = sorted(stock.get(p.sku, []), key=lambda r: (r[0] != p.location_id, -r[1], r[0]))
        passing = sum(q for loc, q in places if loc == DISPATCH)
        out_let = sum(q for loc, q in places if loc in outlet)
        bodega = sum(q for loc, q in places if loc != DISPATCH and loc not in outlet)
        reserve = sum(it.qty for it in inv._reserve_for(p, index))
        status = "Agotado" if bodega <= 0 else "Bajo mínimo" if p.min_qty > 0 and bodega <= p.min_qty else "OK"
        order = max(p.min_qty * 2 - bodega - reserve, 0) if p.min_qty > 0 and bodega <= p.min_qty else 0
        bring = min(reserve, max(p.min_qty * 2 - bodega, 0)) if p.min_qty > 0 and bodega <= p.min_qty else 0
        info = cat.get(p.sku) or {}
        price = info.get("price") or 0
        rows.append({
            "p": p, "bodega": bodega, "reserve": reserve, "passing": passing, "outlet": out_let,
            "total": bodega + reserve, "status": status, "order": order, "bring": bring,
            "main": names.get(p.location_id, p.location_id),
            "where": " · ".join(f"{loc} ({q})" for loc, q in places),
            "places": places, "out30": outs.get(p.sku, 0), "price": price or None,
            "value": bodega * price if price else None, "url": _store_url(info.get("product") or ""),
        })
    return {"rows": rows, "names": names, "outlet": outlet}


# ------------------------------------------------------------------ libros
def inventory_workbook(db: Session, who: str) -> bytes:
    data = _inventory_rows(db)
    rows = data["rows"]
    now = datetime.now(ser.BOGOTA)
    stamp = f"Generado el {_when(now)} · por {who}"
    wb = Workbook()
    wb.properties.title = "Inventario · Bodega Pigmalion"
    wb.properties.creator = "Bodega"

    # --- Inventario (una fila por talla) ---
    ws = wb.active
    ws.title = "Inventario"
    headers = ["Referencia", "Talla", "Código", "Bodega", "Reserva", "De paso", "Outlet", "Total", "Mínimo", "Estado",
               "Pedir", "Ubicación principal", "Dónde está", "Salidas 30 días", "Precio tienda", "Valor bodega", "Tienda"]
    body = [[r["p"].name, r["p"].size or "Única", r["p"].sku, r["bodega"], r["reserve"], r["passing"], r["outlet"],
             r["total"], r["p"].min_qty, r["status"], r["order"], r["main"], r["where"], r["out30"], r["price"], r["value"],
             "Ver" if r["url"] else ""] for r in rows]
    # sin la tienda (sin internet o sin catalogo): sin las columnas de precio, valor y enlace
    drop = ({"Precio tienda", "Valor bodega"} if not any(r["price"] for r in rows) else set()) | (
        {"Tienda"} if not any(r["url"] for r in rows) else set())
    keep = [i for i, h in enumerate(headers) if h not in drop]
    headers, body = [headers[i] for i in keep], [[row[i] for i in keep] for row in body]
    col = lambda h: get_column_letter(headers.index(h) + 1)
    _title(ws, "Inventario por talla", f"{stamp} · Bodega = en los muebles, sin lo de paso ni el outlet", len(headers))
    fmt = {h: INT for h in ("Bodega", "Reserva", "De paso", "Outlet", "Total", "Mínimo", "Pedir", "Salidas 30 días")}
    fmt.update({"Precio tienda": MONEY, "Valor bodega": MONEY})
    last = _table(ws, "Inventario", 4, headers, body, fmt, band_by=0,
                  totals=[h for h in ("Bodega", "Reserva", "De paso", "Outlet", "Total", "Pedir", "Salidas 30 días", "Valor bodega") if h in headers])
    for i, r in enumerate(rows, 5):  # el enlace a la tienda
        if r["url"] and "Tienda" in headers:
            cell = ws.cell(row=i, column=headers.index("Tienda") + 1, value="Ver")
            cell.hyperlink = r["url"]
            cell.font = Font(color="1D4ED8", underline="single")
        for c in range(1, 4):
            ws.cell(row=i, column=c).font = Font(bold=c == 1)
        ws.cell(row=i, column=3).font = Font(name="Consolas")
        ws.cell(row=i, column=4).font = Font(bold=True)
    _status_colors(ws, col("Estado"), 5, last)
    _bars(ws, col("Salidas 30 días"), 5, last)
    if "Tienda" in headers:
        ws.column_dimensions[col("Tienda")].width = 8
    ws.freeze_panes = "D5"
    _print(ws, 4)

    # --- Por referencia ---
    wr = wb.create_sheet("Por referencia")
    by_ref: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        by_ref[r["p"].name].append(r)
    rh = ["Referencia", "Tallas en la bodega", "Bodega", "Reserva", "Total", "Tallas agotadas", "Tallas bajo mínimo",
          "Pedir", "Salidas 30 días", "Ubicaciones"]
    rbody = []
    for name, rs in by_ref.items():
        locs = sorted({loc for r in rs for loc, _ in r["places"]})
        rbody.append([name, " · ".join(f"{r['p'].size or 'Única'} {r['bodega']}" for r in rs),
                      sum(r["bodega"] for r in rs), sum(r["reserve"] for r in rs), sum(r["total"] for r in rs),
                      sum(1 for r in rs if r["status"] == "Agotado"), sum(1 for r in rs if r["status"] == "Bajo mínimo"),
                      sum(r["order"] for r in rs), sum(r["out30"] for r in rs), ", ".join(locs)])
    _title(wr, "Inventario por referencia", stamp, len(rh))
    last = _table(wr, "PorReferencia", 4, rh, rbody,
                  {h: INT for h in ("Bodega", "Reserva", "Total", "Tallas agotadas", "Tallas bajo mínimo", "Pedir", "Salidas 30 días")},
                  totals=["Bodega", "Reserva", "Total", "Pedir", "Salidas 30 días"])
    for i in range(5, last + 1):
        wr.cell(row=i, column=1).font = Font(bold=True)
    _bars(wr, "C", 5, last)
    wr.freeze_panes = "B5"
    _print(wr, 4)

    # --- Por ubicacion (para contar mueble por mueble) ---
    wl = wb.create_sheet("Por ubicación")
    lh = ["Ubicación", "Código ubicación", "Tipo", "Referencia", "Talla", "Código", "Cantidad"]
    lbody = []
    for r in rows:
        for loc, q in r["places"]:
            kind = "De paso" if loc == DISPATCH else "Outlet" if loc in data["outlet"] else "Bodega"
            lbody.append([data["names"].get(loc, loc), loc, kind, r["p"].name, r["p"].size or "Única", r["p"].sku, q])
    natural = lambda s: [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", s)]
    lbody.sort(key=lambda x: (x[2] != "Bodega", natural(x[1]), x[3], size_rank(x[4])))
    _title(wl, "Inventario por ubicación", f"{stamp} · para contar mueble por mueble", len(lh))
    last = _table(wl, "PorUbicacion", 4, lh, lbody, {"Cantidad": INT}, band_by=1, totals=["Cantidad"])
    for i in range(5, last + 1):
        wl.cell(row=i, column=2).font = Font(name="Consolas", bold=True)
    wl.freeze_panes = "D5"
    _print(wl, 4, landscape=False)

    # --- Reserva ---
    wres = wb.create_sheet("Reserva")
    in_bodega = {r["p"].sku: r["bodega"] for r in rows}
    by_name = {inv.ref_key(r["p"].name, r["p"].size): r["bodega"] for r in rows}
    items = sorted(db.query(models.ReserveItem).filter(models.ReserveItem.qty > 0).all(),
                   key=lambda it: (it.name, size_rank(it.size)))
    resh = ["Referencia", "Talla", "Código", "En la reserva", "En la bodega hay", "Nota"]
    resbody = []
    for it in items:
        have = in_bodega.get(it.sku) if it.sku else by_name.get(inv.ref_key(it.name, it.size))
        note = ("Sin código: ponle el código al llevarla a la bodega" if not it.sku
                else "Agotada en la bodega: llévala" if have == 0 else "No está registrada en la bodega" if have is None else "")
        resbody.append([it.name, it.size or "Única", it.sku or "", it.qty, have, note])
    _title(wres, "Bodega de reserva", stamp, len(resh))
    last = _table(wres, "Reserva", 4, resh, resbody, {"En la reserva": INT, "En la bodega hay": INT}, band_by=0,
                  totals=["En la reserva"])
    wres.freeze_panes = "B5"
    _print(wres, 4, landscape=False)

    # --- Por pedir ---
    _needs_sheet(wb.create_sheet("Por pedir"), rows, stamp)

    # --- Resumen (primera hoja) ---
    wsum = wb.create_sheet("Resumen", 0)
    _summary(wsum, rows, by_ref, items, stamp)
    wb.active = 0
    return _save(wb)


def _needs_sheet(ws, rows: list[dict], stamp: str) -> None:
    need = [r for r in rows if r["p"].min_qty > 0 and r["bodega"] <= r["p"].min_qty]
    need.sort(key=lambda r: (r["bodega"] / r["p"].min_qty, r["p"].name, size_rank(r["p"].size)))
    h = ["Referencia", "Talla", "Código", "Hay en la bodega", "Mínimo", "En la reserva", "Traer de la reserva", "Pedir",
         "Estado", "Ubicación principal"]
    body = [[r["p"].name, r["p"].size or "Única", r["p"].sku, r["bodega"], r["p"].min_qty, r["reserve"], r["bring"],
             r["order"], r["status"], r["main"]] for r in need]
    _title(ws, "Por pedir", f"{stamp} · se pide para llegar al doble del mínimo; lo que hay en la reserva se trae, no se compra", len(h))
    last = _table(ws, "PorPedir", 4, h, body, {x: INT for x in ("Hay en la bodega", "Mínimo", "En la reserva", "Traer de la reserva", "Pedir")},
                  totals=["Traer de la reserva", "Pedir"])
    for i in range(5, last + 1):
        ws.cell(row=i, column=8).font = Font(bold=True)
        ws.cell(row=i, column=3).font = Font(name="Consolas")
    _status_colors(ws, "I", 5, last)
    ws.freeze_panes = "D5"
    _print(ws, 4)


def _summary(ws, rows: list[dict], by_ref: dict, reserve_items: list, stamp: str) -> None:
    _title(ws, "Inventario · Bodega Pigmalion", stamp, 4)
    priced = [r for r in rows if r["price"]]
    kpis = [
        ("Referencias", len(by_ref), INT),
        ("Códigos (tallas)", len(rows), INT),
        ("Prendas en la bodega", sum(r["bodega"] for r in rows), INT),
        ("En la reserva", sum(it.qty for it in reserve_items), INT),
        ("De paso (para despachar)", sum(r["passing"] for r in rows), INT),
        ("En outlet", sum(r["outlet"] for r in rows), INT),
        ("Tallas agotadas", sum(1 for r in rows if r["status"] == "Agotado"), INT),
        ("Tallas bajo mínimo", sum(1 for r in rows if r["status"] == "Bajo mínimo"), INT),
        ("Prendas por pedir", sum(r["order"] for r in rows), INT),
        ("Prendas para traer de la reserva", sum(r["bring"] for r in rows), INT),
        ("Salidas de los últimos 30 días", sum(r["out30"] for r in rows), INT),
    ]
    if priced:
        kpis.append(("Valor de la bodega (precio de la tienda)", sum(r["value"] or 0 for r in rows), MONEY))
    ws["A4"] = "En números"
    ws["A4"].font = Font(size=12, bold=True, color=INK)
    for i, (label, value, fmt) in enumerate(kpis, 5):
        a, b = ws.cell(row=i, column=1, value=label), ws.cell(row=i, column=2, value=value)
        b.number_format = fmt
        b.font = Font(size=13, bold=True)
        b.alignment = Alignment(horizontal="right")
        for c in (a, b):
            c.fill = PatternFill("solid", fgColor=FOG if i % 2 else WHITE)
            c.border = Border(bottom=Side(style="thin", color="E4E8E0"))
    ws.cell(row=5, column=2).fill = PatternFill("solid", fgColor=LIME)  # la primera cifra, en verde Pigmalion
    top = sorted([r for r in rows if r["out30"]], key=lambda r: -r["out30"])[:10]
    ws["D4"] = "Lo que más sale (30 días)"
    ws["D4"].font = Font(size=12, bold=True, color=INK)
    if top:
        th = ["Prenda", "Salieron", "Hay"]
        for c, h in enumerate(th, 4):
            cell = ws.cell(row=5, column=c, value=h)
            cell.font, cell.fill = Font(bold=True, color=WHITE), PatternFill("solid", fgColor=INK)
        for i, r in enumerate(top, 6):
            ws.cell(row=i, column=4, value=f"{r['p'].name} · {r['p'].size or 'Única'}")
            ws.cell(row=i, column=5, value=r["out30"]).number_format = INT
            ws.cell(row=i, column=6, value=r["bodega"]).number_format = INT
        _bars(ws, "E", 6, 5 + len(top))
    else:
        ws["D5"] = "Todavía no hay salidas en los últimos 30 días."
        ws["D5"].font = Font(italic=True, color=MUTED)
    note_row = max(6 + len(kpis), 7 + len(top)) + 1
    ws.cell(row=note_row, column=1, value="Cómo leer este archivo").font = Font(size=12, bold=True)
    notes = [
        "Inventario: una fila por talla. Bodega es lo que hay en los muebles, sin lo de paso ni el outlet.",
        "Cada encabezado tiene filtro: por ejemplo, Estado = Bajo mínimo. El total de abajo suma solo lo que se ve.",
        "Por referencia: cada referencia con sus tallas. Por ubicación: para contar mueble por mueble.",
        "Por pedir: lo que está en su mínimo o agotado. Primero se trae de la reserva y el resto se pide.",
    ]
    for i, n in enumerate(notes, note_row + 1):
        ws.cell(row=i, column=1, value=f"• {n}").font = Font(color=MUTED)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 46
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 10
    ws.sheet_view.showGridLines = False
    _print(ws, 4, landscape=False)


def needs_workbook(db: Session, who: str) -> bytes:
    """Solo lo que hay que pedir (el boton "Descargar" de Por reponer)."""
    rows = _inventory_rows(db)["rows"]
    wb = Workbook()
    wb.properties.title = "Por pedir · Bodega Pigmalion"
    ws = wb.active
    ws.title = "Por pedir"
    _needs_sheet(ws, rows, f"Generado el {_when(datetime.now(ser.BOGOTA))} · por {who}")
    return _save(wb)


TYPE = {"in": "Entrada", "out": "Salida", "set": "Conteo", "new": "Registro nuevo", "move": "Traslado"}


def movements_workbook(db: Session, who: str, limit: int = 5000) -> bytes:
    names = ser.loc_names(db)
    moves = db.query(models.Movement).order_by(models.Movement.id.desc()).limit(limit).all()
    wb = Workbook()
    wb.properties.title = "Historial · Bodega Pigmalion"
    ws = wb.active
    ws.title = "Historial"
    h = ["Fecha", "Tipo", "Referencia", "Talla", "Código", "Cambio", "Antes", "Después", "Ubicación", "Hacia", "Nota", "Usuario"]
    body = []
    per_day: dict = defaultdict(lambda: [0, 0, 0])
    for m in moves:
        when = ser.local_time(m.created_at).replace(tzinfo=None)
        change = m.after - m.before if m.type in ("set", "move") else (m.qty if m.type in ("in", "new") else -m.qty)
        if m.type == "move":
            change = 0
        body.append([when, TYPE.get(m.type, m.type), m.product_name, m.product_size or "Única", m.sku, change, m.before,
                     m.after, names.get(m.location_id, m.location_id),
                     names.get(m.to_location_id, m.to_location_id) if m.to_location_id else "", m.note or "", m.user_name])
        d = per_day[when.date()]
        if m.type in ("in", "new"):
            d[0] += m.qty
        elif m.type == "out":
            d[1] += m.qty
        elif m.type == "set":
            d[2] += abs(m.after - m.before)
    _title(ws, "Historial de movimientos", f"Generado el {_when(datetime.now(ser.BOGOTA))} · por {who} · los últimos {len(body)}", len(h))
    last = _table(ws, "Historial", 4, h, body, {"Fecha": DATE, "Cambio": '+#,##0;-#,##0;0', "Antes": INT, "Después": INT})
    ws.column_dimensions["A"].width = 17
    for i in range(5, last + 1):
        ws.cell(row=i, column=5).font = Font(name="Consolas")
    green, red = PatternFill("solid", fgColor="E3F7C2"), PatternFill("solid", fgColor="FDE2E1")
    ws.conditional_formatting.add(f"F5:F{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=green, font=Font(color="3B5B00", bold=True)))
    ws.conditional_formatting.add(f"F5:F{last}", CellIsRule(operator="lessThan", formula=["0"], fill=red, font=Font(color="B42318", bold=True)))
    ws.freeze_panes = "C5"
    _print(ws, 4)

    wd = wb.create_sheet("Por día")
    dh = ["Día", "Entraron", "Salieron", "Ajustes de conteo", "Neto"]
    dbody = [[day, a, b, c, a - b] for day, (a, b, c) in sorted(per_day.items(), reverse=True)]
    _title(wd, "Entradas y salidas por día", "Entraron = entradas y registros nuevos · Salieron = salidas · Neto = entraron − salieron", len(dh))
    last = _table(wd, "PorDia", 4, dh, dbody, {"Día": "dd/mm/yyyy", "Entraron": INT, "Salieron": INT,
                                               "Ajustes de conteo": INT, "Neto": '+#,##0;-#,##0;0'},
                  totals=["Entraron", "Salieron", "Ajustes de conteo", "Neto"])
    _bars(wd, "B", 5, last, LIME_DEEP)
    _bars(wd, "C", 5, last, INK)
    wd.column_dimensions["A"].width = 13
    wd.freeze_panes = "B5"
    _print(wd, 4, landscape=False)
    return _save(wb)
