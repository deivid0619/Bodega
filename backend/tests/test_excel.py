"""Los Excel que se descargan: archivos .xlsx de verdad, con sus hojas, la
tabla con filtros y lo de cada prenda separado (bodega, reserva...)."""
import io
import os

os.environ.setdefault("DATABASE_URL", "sqlite:///./test_bodega.db")
os.environ.setdefault("ADMIN_EMAIL", "admin@test.com")
os.environ.setdefault("ADMIN_PASSWORD", "admin1234")

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.config import settings
from app.main import app


def _h(token):
    return {"Authorization": f"Bearer {token}"}


def test_inventory_movements_and_order_in_excel():
    with TestClient(app) as client:
        admin = _h(client.post("/api/auth/login", json={"email": settings.admin_email, "password": settings.admin_password}).json()["access_token"])
        r = client.post("/api/products", headers=admin, json={"sku": "XLS-M", "name": "PRUEBA EXCEL", "size": "M",
                                                              "location_id": "F-1-1", "qty": 2, "min_qty": 3})
        assert r.status_code in (201, 400)  # si quedo de otra prueba, igual sirve
        client.post("/api/reserve", headers=admin, json={"name": "PRUEBA EXCEL", "size": "M", "sku": "XLS-M", "qty": 5})

        r = client.get("/api/reports/inventory.xlsx", headers=admin)
        assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
        assert r.headers["content-disposition"].endswith('.xlsx"')
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Resumen", "Inventario", "Por referencia", "Por ubicación", "Reserva", "Por pedir"]
        ws = wb["Inventario"]
        assert [c.value for c in ws[4]][:10] == ["Referencia", "Talla", "Código", "Bodega", "Reserva", "De paso", "Outlet",
                                                 "Total", "Mínimo", "Estado"]
        assert "Inventario" in ws.tables and ws.freeze_panes == "D5"
        row = next(r for r in ws.iter_rows(min_row=5, values_only=True) if r[2] == "XLS-M")
        bodega, reserve = row[3], row[4]
        assert reserve >= 5 and row[7] == bodega + reserve  # el total suma bodega y reserva
        assert row[9] in ("Bajo mínimo", "Agotado")  # en su minimo o menos
        assert any(str(c.value or "").startswith("=SUBTOTAL(109") for c in ws[ws.max_row])  # totales con el filtro
        pedir = wb["Por pedir"]
        assert any(r[2] == "XLS-M" for r in pedir.iter_rows(min_row=5, values_only=True))

        r = client.get("/api/reports/movements.xlsx", headers=admin)
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Historial", "Por día"]
        assert [c.value for c in wb["Historial"][4]][:6] == ["Fecha", "Tipo", "Referencia", "Talla", "Código", "Cambio"]

        r = client.get("/api/reports/pedido.xlsx", headers=admin)
        assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Por pedir"]
